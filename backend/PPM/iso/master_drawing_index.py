"""
FastAPI Router & Document Generator for ISO Document: Master Drawing Index.
Generates Word (.docx) document matching CMTI-QMS Master Drawing Index format.
"""

from typing import List, Dict, Any, Optional
import io
from fastapi import APIRouter, HTTPException, status, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from iso.header import add_header_table
from iso.finalfooter import add_footer_table
from iso.sqap import set_cell_shading, set_cell_border, set_cell_margins, add_text

router = APIRouter(prefix="/iso", tags=["ISO Master Drawing Index"])


# ============================================================
# REQUEST MODELS
# ============================================================

class MasterDrawingItemRequest(BaseModel):
    sl_no: str = ""
    assembly_no: str = ""
    size: str = ""
    material: str = ""
    qty: str = ""
    name_of_part: str = ""
    revision: str = ""


class MasterDrawingIndexRequest(BaseModel):
    project_title: str = ""
    project_no: str = ""
    customer_name: str = ""
    sub_system: str = ""
    items: Optional[List[Any]] = None
    prepared_by: str = ""
    approved_by: str = ""
    group_name: str = ""
    centre_dept: str = ""
    doc_no: str = "066"
    doc_date: str = ""
    filename: str = "ISO_Master_Drawing_Index.docx"


DEFAULT_DRAWING_INDEX_HEADERS = [
    "Assembly No. /\nPart No.",
    "Size",
    "Material",
    "Qty",
    "Name of Part/Assembly",
    "Revision",
    "Sl."
]


# ============================================================
# DOCUMENT GENERATOR
# ============================================================

def create_master_drawing_index_document(
    project_title: str = "",
    project_no: str = "",
    customer_name: str = "",
    sub_system: str = "",
    items: Optional[List[Any]] = None,
    prepared_by: str = "",
    approved_by: str = "",
    group_name: str = "",
    centre_dept: str = "",
    doc_no: str = "066",
    doc_date: str = ""
) -> Document:
    doc = Document()

    # Margins (A4 Landscape, 0.5" all sides)
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)

    header_group = (group_name or centre_dept or "SMC").strip().upper()
    if header_group.startswith("G-") or header_group.startswith("C-"):
        header_group = header_group[2:]

    # Add Standard ISO Header Table
    add_header_table(
        doc.sections[0],
        title=f"MASTER DRAWING INDEX-{header_group}" if header_group else "MASTER DRAWING INDEX",
        page_str="1 of 1",
        centre_dept=centre_dept,
        doc_no=doc_no or "066",
        date_str=doc_date
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(3)

    border_fmt = {"val": "single", "sz": "4", "color": "000000"}
    hdr_bg = "D9E2EC"

    # Metadata Card (Project Title, Project No, Customer Name, Sub-System / Module)
    if project_title or project_no or customer_name or sub_system:
        meta_table = doc.add_table(rows=1, cols=4)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_table.autofit = False

        col_w = Inches(10.69 / 4.0)
        row = meta_table.rows[0]
        meta_data = [
            ("Project Title:", project_title or "--"),
            ("Project No:", project_no or "--"),
            ("Customer Name:", customer_name or "--"),
            ("Sub-system / Module:", sub_system or "--")
        ]

        for c_idx, (lbl, val) in enumerate(meta_data):
            cell = row.cells[c_idx]
            cell.width = col_w
            set_cell_border(cell, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
            set_cell_margins(cell, top=20, start=20, bottom=20, end=20)
            set_cell_shading(cell, "F8FAFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_text(p, f"{lbl} ", font_size=8.5, bold=True)
            add_text(p, str(val), font_size=8.5, bold=False)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # Master Drawing Index Table
    col_widths = [
        Inches(2.0),   # Assembly No. / Part No.
        Inches(0.9),   # Size
        Inches(1.8),   # Material
        Inches(0.8),   # Qty
        Inches(3.4),   # Name of Part/Assembly
        Inches(1.0),   # Revision
        Inches(0.79)   # Sl.
    ]

    row_data_list = items if items else []
    total_rows = 1 + max(1, len(row_data_list))

    table = doc.add_table(rows=total_rows, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_row = table.rows[0]
    for c_idx, h_text in enumerate(DEFAULT_DRAWING_INDEX_HEADERS):
        cell = hdr_row.cells[c_idx]
        cell.width = col_widths[c_idx]
        set_cell_border(cell, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
        set_cell_shading(cell, hdr_bg)
        set_cell_margins(cell, top=25, start=20, bottom=25, end=20)
        add_text(cell, h_text, font_size=9, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    if not row_data_list:
        blank_row = table.rows[1]
        for c_idx in range(7):
            cell = blank_row.cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_border(cell, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
            set_cell_margins(cell, top=20, start=20, bottom=20, end=20)
            add_text(cell, "", font_size=9)
    else:
        for r_idx, item in enumerate(row_data_list):
            row = table.rows[r_idx + 1]
            if isinstance(item, dict):
                assm_no = item.get("assembly_no") or item.get("part_no") or ""
                size = item.get("size") or ""
                mat = item.get("material") or ""
                qty = str(item.get("qty") or "")
                name_part = item.get("name_of_part") or item.get("part_name") or ""
                rev = item.get("revision") or item.get("rev") or ""
                sl = item.get("sl_no") or str(r_idx + 1)
            elif isinstance(item, list):
                assm_no = item[0] if len(item) > 0 else ""
                size = item[1] if len(item) > 1 else ""
                mat = item[2] if len(item) > 2 else ""
                qty = str(item[3]) if len(item) > 3 else ""
                name_part = item[4] if len(item) > 4 else ""
                rev = item[5] if len(item) > 5 else ""
                sl = str(item[6]) if len(item) > 6 else str(r_idx + 1)
            else:
                assm_no = getattr(item, "assembly_no", "") or getattr(item, "part_no", "")
                size = getattr(item, "size", "")
                mat = getattr(item, "material", "")
                qty = str(getattr(item, "qty", ""))
                name_part = getattr(item, "name_of_part", "") or getattr(item, "part_name", "")
                rev = getattr(item, "revision", "") or getattr(item, "rev", "")
                sl = getattr(item, "sl_no", "") or str(r_idx + 1)

            vals = [assm_no, size, mat, qty, name_part, rev, sl]

            for c_idx, val in enumerate(vals):
                cell = row.cells[c_idx]
                cell.width = col_widths[c_idx]
                set_cell_border(cell, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
                set_cell_margins(cell, top=20, start=20, bottom=20, end=20)
                # Align Center for Size, Qty, Revision, Sl; Left for Assembly No, Material, Name of Part
                align = WD_ALIGN_PARAGRAPH.CENTER if c_idx in (1, 3, 5, 6) else WD_ALIGN_PARAGRAPH.LEFT
                add_text(cell, str(val or ""), font_size=9, alignment=align)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Standard ISO Footer Table
    doc_code_footer = f"CMTI-{header_group}-QMS-{doc_no or '066'}/Rev00"
    add_footer_table(
        doc,
        prepared_name=prepared_by,
        approved_name=approved_by,
        group_name=header_group,
        doc_code=doc_code_footer,
        in_body=True
    )

    return doc


def generate_master_drawing_index_bytes(
    project_title: str = "",
    project_no: str = "",
    customer_name: str = "",
    sub_system: str = "",
    items: Optional[List[Any]] = None,
    prepared_by: str = "",
    approved_by: str = "",
    group_name: str = "",
    centre_dept: str = "",
    doc_no: str = "066",
    doc_date: str = ""
) -> bytes:
    doc = create_master_drawing_index_document(
        project_title=project_title,
        project_no=project_no,
        customer_name=customer_name,
        sub_system=sub_system,
        items=items,
        prepared_by=prepared_by,
        approved_by=approved_by,
        group_name=group_name,
        centre_dept=centre_dept,
        doc_no=doc_no,
        doc_date=doc_date
    )
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


# ============================================================
# ROUTER ENDPOINTS
# ============================================================

@router.post("/master-drawing-index/generate")
def generate_master_drawing_index_post(payload: MasterDrawingIndexRequest):
    try:
        data = generate_master_drawing_index_bytes(
            project_title=payload.project_title,
            project_no=payload.project_no,
            customer_name=payload.customer_name,
            sub_system=payload.sub_system,
            items=payload.items,
            prepared_by=payload.prepared_by,
            approved_by=payload.approved_by,
            group_name=payload.group_name,
            centre_dept=payload.centre_dept,
            doc_no=payload.doc_no,
            doc_date=payload.doc_date
        )
        fn = payload.filename or f"ISO_Master_Drawing_Index_{payload.project_no or '066'}.docx"
        if not fn.endswith(".docx"):
            fn += ".docx"
        return StreamingResponse(
            io.BytesIO(data),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f'attachment; filename="{fn}"'}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Master Drawing Index generation failed: {str(e)}")


@router.get("/master-drawing-index/generate")
def generate_master_drawing_index_get(
    project_title: str = Query("", description="Project Title"),
    project_no: str = Query("", description="Project Number"),
    customer_name: str = Query("", description="Customer Name"),
    sub_system: str = Query("", description="Sub-system / Module"),
    prepared_by: str = Query("", description="Prepared By"),
    approved_by: str = Query("", description="Approved By"),
    group_name: str = Query("SMC", description="Group Name"),
    centre_dept: str = Query("", description="Centre / Dept"),
    doc_no: str = Query("066", description="Document No"),
    doc_date: str = Query("", description="Document Date"),
    filename: str = Query("ISO_Master_Drawing_Index.docx")
):
    try:
        data = generate_master_drawing_index_bytes(
            project_title=project_title,
            project_no=project_no,
            customer_name=customer_name,
            sub_system=sub_system,
            items=[],
            prepared_by=prepared_by,
            approved_by=approved_by,
            group_name=group_name,
            centre_dept=centre_dept,
            doc_no=doc_no,
            doc_date=doc_date
        )
        return StreamingResponse(
            io.BytesIO(data),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Master Drawing Index generation failed: {str(e)}")
