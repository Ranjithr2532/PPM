"""
FastAPI Router & Document Generator for ISO Document 065: Technical Specification Format.
Generates Word (.docx) document matching CMTI-QMS-065/Rev00 specification.
"""

import io
from typing import List, Dict, Any, Optional

from fastapi import APIRouter, status, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

from iso.header import add_header_table
from iso.finalfooter import add_footer_table
from iso.sqap import set_cell_shading, set_cell_border, set_cell_margins, add_text

router = APIRouter(prefix="/iso", tags=["ISO Technical Specification (Doc 065)"])


# ============================================================
# REQUEST MODELS
# ============================================================

class SpecItemRequest(BaseModel):
    sl_no: str = ""
    specification: str = ""
    requirement: str = ""
    vendor_compliance: str = ""


class ScopeItemRequest(BaseModel):
    sl_no: str = ""
    particulars: str = ""
    qty: str = ""
    remarks: str = ""


class ChecklistItemRequest(BaseModel):
    sl_no: str = ""
    document_type: str = ""
    checked: bool = False
    remarks: str = ""


class TechnicalSpecificationRequest(BaseModel):
    project_title: str = ""
    project_no: str = ""
    customer_name: str = ""
    item_description: str = ""
    specs: Optional[List[Any]] = None
    scope_of_supply: Optional[List[Any]] = None
    boi_checklist: Optional[List[Any]] = None
    mfg_checklist: Optional[List[Any]] = None
    prepared_by: str = ""
    approved_by: str = ""
    group_name: str = ""
    centre_dept: str = ""
    doc_no: str = "065"
    doc_date: str = ""
    filename: str = "ISO_Technical_Specification.docx"


DEFAULT_BOI_CHECKLIST = [
    {"sl_no": "1", "document_type": "ISO 9001-2015 / QMS Certificate", "checked": True, "remarks": ""},
    {"sl_no": "2", "document_type": "Calibration Certificate", "checked": True, "remarks": ""},
    {"sl_no": "3", "document_type": "Certificate of Compliance (COC)", "checked": True, "remarks": ""},
    {"sl_no": "4", "document_type": "Warranty Certificate", "checked": True, "remarks": ""},
    {"sl_no": "5", "document_type": "CE/ Relevant Standard (For Electrical)", "checked": False, "remarks": ""},
    {"sl_no": "6", "document_type": "User Manual", "checked": True, "remarks": ""},
]

DEFAULT_MFG_CHECKLIST = [
    {"sl_no": "1", "document_type": "ISO 9001-2015 / QMS Certificate", "checked": True, "remarks": ""},
    {"sl_no": "2", "document_type": "Material Chemical Certificate", "checked": True, "remarks": ""},
    {"sl_no": "3", "document_type": "NDT Test Certificate (Ultrasonic Test, DPT, MPT)", "checked": False, "remarks": ""},
    {"sl_no": "4", "document_type": "Heat Treatment Certificate (Hardening, Carburizing, Blackening, Tempering, Annealing)", "checked": False, "remarks": ""},
    {"sl_no": "5", "document_type": "QAP Documents As per CMTI Standards (Casting, Forging, Welding, Fabrication, Machining, Gear Cutting)", "checked": True, "remarks": ""},
    {"sl_no": "6", "document_type": "Dimensional Inspection Report", "checked": True, "remarks": ""},
    {"sl_no": "7", "document_type": "Job Card", "checked": False, "remarks": ""},
    {"sl_no": "8", "document_type": "Witness", "checked": False, "remarks": ""},
]


# ============================================================
# DOCUMENT GENERATOR
# ============================================================

def create_technical_specification_document(
    project_title: str = "",
    project_no: str = "",
    customer_name: str = "",
    item_description: str = "",
    specs: Optional[List[Any]] = None,
    scope_of_supply: Optional[List[Any]] = None,
    boi_checklist: Optional[List[Any]] = None,
    mfg_checklist: Optional[List[Any]] = None,
    prepared_by: str = "",
    approved_by: str = "",
    group_name: str = "",
    centre_dept: str = "",
    doc_no: str = "065",
    doc_date: str = "",
) -> Document:
    doc = Document()

    # Page Margins (A4 Portrait, 0.5" all sides)
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    header_group = (group_name or centre_dept or "SMC").strip().upper()
    if header_group.startswith("G-") or header_group.startswith("C-"):
        header_group = header_group[2:]

    # Add Standard ISO Header Table
    add_header_table(
        section,
        title=f"TECHNICAL SPECIFICATION-{header_group}",
        page_str="1 of 1",
        centre_dept=centre_dept,
        doc_no=doc_no or "065",
        date_str=doc_date
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    border_fmt = {"val": "single", "sz": "4", "color": "000000"}
    hdr_bg = "D9E2EC"

    # --- Project Metadata Card ---
    if project_title or project_no or customer_name:
        meta_table = doc.add_table(rows=1, cols=3)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_table.autofit = False

        col_w = Inches(7.27 / 3.0)
        row = meta_table.rows[0]
        meta_data = [
            ("Project Title:", project_title or "--"),
            ("Project No:", project_no or "--"),
            ("Customer Name:", customer_name or "--")
        ]

        for c_idx, (lbl, val) in enumerate(meta_data):
            cell = row.cells[c_idx]
            cell.width = col_w
            set_cell_border(cell, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
            set_cell_margins(cell, top=20, start=20, bottom=20, end=20)
            set_cell_shading(cell, "F8FAFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_text(p, f"{lbl} ", font_size=8.5, bold=True)
            add_text(p, str(val), font_size=8.5, bold=False)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # --- Section 1: Item Description ---
    p_desc = doc.add_paragraph()
    add_text(p_desc, "DESCRIPTION OF ITEM: ", font_size=10, bold=True)
    if item_description:
        add_text(p_desc, item_description, font_size=10, bold=False)
    else:
        add_text(p_desc, "", font_size=10, bold=False)
    p_desc.paragraph_format.space_after = Pt(4)

    # --- Section 1 Table: Technical Specifications ---
    spec_rows = specs if specs else []
    total_spec_rows = 1 + max(1, len(spec_rows))
    tbl_specs = doc.add_table(rows=total_spec_rows, cols=4)
    tbl_specs.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_specs.autofit = False

    spec_widths = [Inches(0.6), Inches(2.2), Inches(2.2), Inches(2.27)]
    spec_headers = ["Sl No", "Specification", "Requirement", "Vendor Compliance"]

    hdr_row = tbl_specs.rows[0]
    for c_idx, h_text in enumerate(spec_headers):
        cell = hdr_row.cells[c_idx]
        cell.width = spec_widths[c_idx]
        set_cell_border(cell, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
        set_cell_shading(cell, hdr_bg)
        set_cell_margins(cell, top=25, start=20, bottom=25, end=20)
        add_text(cell, h_text, font_size=9, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    if not spec_rows:
        row = tbl_specs.rows[1]
        for c_idx in range(4):
            cell = row.cells[c_idx]
            cell.width = spec_widths[c_idx]
            set_cell_border(cell, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
            set_cell_margins(cell, top=20, start=20, bottom=20, end=20)
            add_text(cell, "", font_size=9)
    else:
        for r_idx, item in enumerate(spec_rows):
            row = tbl_specs.rows[r_idx + 1]
            if isinstance(item, dict):
                sl = item.get("sl_no") or str(r_idx + 1)
                sp = item.get("specification", "")
                rq = item.get("requirement", "")
                vc = item.get("vendor_compliance", "")
            elif isinstance(item, list):
                sl = item[0] if len(item) > 0 else str(r_idx + 1)
                sp = item[1] if len(item) > 1 else ""
                rq = item[2] if len(item) > 2 else ""
                vc = item[3] if len(item) > 3 else ""
            else:
                sl = getattr(item, "sl_no", "") or str(r_idx + 1)
                sp = getattr(item, "specification", "")
                rq = getattr(item, "requirement", "")
                vc = getattr(item, "vendor_compliance", "")

            vals = [sl, sp, rq, vc]
            for c_idx, val in enumerate(vals):
                cell = row.cells[c_idx]
                cell.width = spec_widths[c_idx]
                set_cell_border(cell, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
                set_cell_margins(cell, top=20, start=20, bottom=20, end=20)
                add_text(cell, str(val or ""), font_size=9, alignment=WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- Section 2: Scope of Supply ---
    p_scope_title = doc.add_paragraph()
    add_text(p_scope_title, "Scope of Supply", font_size=10.5, bold=True)
    p_scope_title.paragraph_format.space_after = Pt(2)

    p_scope_sub = doc.add_paragraph()
    add_text(p_scope_sub, "Supplier need to supply below mentioned items,", font_size=9.5, italic=True)
    p_scope_sub.paragraph_format.space_after = Pt(4)

    scope_rows = scope_of_supply if scope_of_supply else []
    total_scope_rows = 1 + max(1, len(scope_rows))
    tbl_scope = doc.add_table(rows=total_scope_rows, cols=4)
    tbl_scope.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_scope.autofit = False

    scope_widths = [Inches(0.6), Inches(4.0), Inches(1.1), Inches(1.57)]
    scope_headers = ["Sl.No.", "Particulars", "Qty", "Remarks"]

    hdr_row = tbl_scope.rows[0]
    for c_idx, h_text in enumerate(scope_headers):
        cell = hdr_row.cells[c_idx]
        cell.width = scope_widths[c_idx]
        set_cell_border(cell, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
        set_cell_shading(cell, hdr_bg)
        set_cell_margins(cell, top=25, start=20, bottom=25, end=20)
        add_text(cell, h_text, font_size=9, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    if not scope_rows:
        row = tbl_scope.rows[1]
        for c_idx in range(4):
            cell = row.cells[c_idx]
            cell.width = scope_widths[c_idx]
            set_cell_border(cell, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
            set_cell_margins(cell, top=20, start=20, bottom=20, end=20)
            add_text(cell, "", font_size=9)
    else:
        for r_idx, item in enumerate(scope_rows):
            row = tbl_scope.rows[r_idx + 1]
            if isinstance(item, dict):
                sl = item.get("sl_no") or str(r_idx + 1)
                part = item.get("particulars", "")
                qty = str(item.get("qty", ""))
                rem = item.get("remarks", "")
            elif isinstance(item, list):
                sl = item[0] if len(item) > 0 else str(r_idx + 1)
                part = item[1] if len(item) > 1 else ""
                qty = item[2] if len(item) > 2 else ""
                rem = item[3] if len(item) > 3 else ""
            else:
                sl = getattr(item, "sl_no", "") or str(r_idx + 1)
                part = getattr(item, "particulars", "")
                qty = str(getattr(item, "qty", ""))
                rem = getattr(item, "remarks", "")

            vals = [sl, part, qty, rem]
            for c_idx, val in enumerate(vals):
                cell = row.cells[c_idx]
                cell.width = scope_widths[c_idx]
                set_cell_border(cell, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
                set_cell_margins(cell, top=20, start=20, bottom=20, end=20)
                add_text(cell, str(val or ""), font_size=9, alignment=WD_ALIGN_PARAGRAPH.CENTER if c_idx in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- Section 3: Checklist for Documents Required ---
    p_chk_title = doc.add_paragraph()
    add_text(p_chk_title, "CHECKLIST FOR THE DOCUMENTS REQUIRED FROM VENDOR / SUPPLIER", font_size=10.5, bold=True)
    p_chk_title.paragraph_format.space_after = Pt(4)

    # 3A: For Bought Out Items (BOI)
    p_boi = doc.add_paragraph()
    add_text(p_boi, "For Bought Out Items (BOI)", font_size=9.5, bold=True, italic=True)
    p_boi.paragraph_format.space_after = Pt(2)

    boi_items = boi_checklist if boi_checklist is not None else DEFAULT_BOI_CHECKLIST
    tbl_boi = doc.add_table(rows=1 + len(boi_items), cols=3)
    tbl_boi.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_boi.autofit = False

    chk_widths = [Inches(0.6), Inches(5.2), Inches(1.47)]
    chk_headers = ["Sl No", "Document Type", "Tick mark"]

    hdr_row = tbl_boi.rows[0]
    for c_idx, h_text in enumerate(chk_headers):
        cell = hdr_row.cells[c_idx]
        cell.width = chk_widths[c_idx]
        set_cell_border(cell, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
        set_cell_shading(cell, hdr_bg)
        set_cell_margins(cell, top=25, start=20, bottom=25, end=20)
        add_text(cell, h_text, font_size=9, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    for r_idx, item in enumerate(boi_items):
        row = tbl_boi.rows[r_idx + 1]
        if isinstance(item, dict):
            sl = item.get("sl_no") or str(r_idx + 1)
            dtype = item.get("document_type", "")
            is_chk = bool(item.get("checked", False))
        else:
            sl = getattr(item, "sl_no", "") or str(r_idx + 1)
            dtype = getattr(item, "document_type", "")
            is_chk = bool(getattr(item, "checked", False))

        tick = "☑" if is_chk else "☐"

        for c_idx, val in enumerate([sl, dtype, tick]):
            cell = row.cells[c_idx]
            cell.width = chk_widths[c_idx]
            set_cell_border(cell, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
            set_cell_margins(cell, top=20, start=20, bottom=20, end=20)
            add_text(
                cell,
                str(val or ""),
                font_size=11 if c_idx == 2 else 9,
                bold=(c_idx == 2 and is_chk),
                alignment=WD_ALIGN_PARAGRAPH.CENTER if c_idx in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT
            )

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 3B: For Manufacturing Items
    p_mfg = doc.add_paragraph()
    add_text(p_mfg, "For Manufacturing Items", font_size=9.5, bold=True, italic=True)
    p_mfg.paragraph_format.space_after = Pt(2)

    mfg_items = mfg_checklist if mfg_checklist is not None else DEFAULT_MFG_CHECKLIST
    tbl_mfg = doc.add_table(rows=1 + len(mfg_items), cols=3)
    tbl_mfg.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_mfg.autofit = False

    hdr_row = tbl_mfg.rows[0]
    for c_idx, h_text in enumerate(chk_headers):
        cell = hdr_row.cells[c_idx]
        cell.width = chk_widths[c_idx]
        set_cell_border(cell, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
        set_cell_shading(cell, hdr_bg)
        set_cell_margins(cell, top=25, start=20, bottom=25, end=20)
        add_text(cell, h_text, font_size=9, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    for r_idx, item in enumerate(mfg_items):
        row = tbl_mfg.rows[r_idx + 1]
        if isinstance(item, dict):
            sl = item.get("sl_no") or str(r_idx + 1)
            dtype = item.get("document_type", "")
            is_chk = bool(item.get("checked", False))
        else:
            sl = getattr(item, "sl_no", "") or str(r_idx + 1)
            dtype = getattr(item, "document_type", "")
            is_chk = bool(getattr(item, "checked", False))

        tick = "☑" if is_chk else "☐"

        for c_idx, val in enumerate([sl, dtype, tick]):
            cell = row.cells[c_idx]
            cell.width = chk_widths[c_idx]
            set_cell_border(cell, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
            set_cell_margins(cell, top=20, start=20, bottom=20, end=20)
            add_text(
                cell,
                str(val or ""),
                font_size=11 if c_idx == 2 else 9,
                bold=(c_idx == 2 and is_chk),
                alignment=WD_ALIGN_PARAGRAPH.CENTER if c_idx in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT
            )

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Standard ISO Footer Table
    doc_code_footer = f"CMTI-{header_group}-QMS-{doc_no or '065'}/Rev00"
    add_footer_table(
        doc,
        prepared_name=prepared_by,
        approved_name=approved_by,
        group_name=header_group,
        doc_code=doc_code_footer,
        in_body=True
    )

    return doc


def generate_technical_specification_bytes(
    project_title: str = "",
    project_no: str = "",
    customer_name: str = "",
    item_description: str = "",
    specs: Optional[List[Any]] = None,
    scope_of_supply: Optional[List[Any]] = None,
    boi_checklist: Optional[List[Any]] = None,
    mfg_checklist: Optional[List[Any]] = None,
    prepared_by: str = "",
    approved_by: str = "",
    group_name: str = "",
    centre_dept: str = "",
    doc_no: str = "065",
    doc_date: str = "",
) -> io.BytesIO:
    doc = create_technical_specification_document(
        project_title=project_title,
        project_no=project_no,
        customer_name=customer_name,
        item_description=item_description,
        specs=specs,
        scope_of_supply=scope_of_supply,
        boi_checklist=boi_checklist,
        mfg_checklist=mfg_checklist,
        prepared_by=prepared_by,
        approved_by=approved_by,
        group_name=group_name,
        centre_dept=centre_dept,
        doc_no=doc_no,
        doc_date=doc_date
    )

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


# ============================================================
# API ROUTE HANDLERS
# ============================================================

@router.get(
    "/technical-specification/generate",
    status_code=status.HTTP_200_OK,
    summary="Generate ISO Technical Specification Document (.docx) via GET"
)
async def generate_technical_specification_doc_get(
    project_title: str = Query(""),
    project_no: str = Query(""),
    customer_name: str = Query(""),
    item_description: str = Query(""),
    prepared_by: str = Query(""),
    approved_by: str = Query(""),
    group_name: str = Query(""),
    centre_dept: str = Query(""),
    doc_no: str = Query("065"),
    doc_date: str = Query(""),
    filename: str = Query("ISO_Technical_Specification.docx")
):
    buffer = generate_technical_specification_bytes(
        project_title=project_title,
        project_no=project_no,
        customer_name=customer_name,
        item_description=item_description,
        prepared_by=prepared_by,
        approved_by=approved_by,
        group_name=group_name,
        centre_dept=centre_dept,
        doc_no=doc_no,
        doc_date=doc_date
    )

    headers = {
        "Content-Disposition": f'attachment; filename="{filename}"'
    }

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers=headers
    )


@router.post(
    "/technical-specification/generate",
    status_code=status.HTTP_200_OK,
    summary="Generate ISO Technical Specification Document (.docx) via POST"
)
async def generate_technical_specification_doc_post(
    payload: TechnicalSpecificationRequest
):
    filename = payload.filename or "ISO_Technical_Specification.docx"
    if not filename.lower().endswith(".docx"):
        filename += ".docx"

    buffer = generate_technical_specification_bytes(
        project_title=payload.project_title,
        project_no=payload.project_no,
        customer_name=payload.customer_name,
        item_description=payload.item_description,
        specs=payload.specs,
        scope_of_supply=payload.scope_of_supply,
        boi_checklist=payload.boi_checklist,
        mfg_checklist=payload.mfg_checklist,
        prepared_by=payload.prepared_by,
        approved_by=payload.approved_by,
        group_name=payload.group_name,
        centre_dept=payload.centre_dept,
        doc_no=payload.doc_no,
        doc_date=payload.doc_date
    )

    headers = {
        "Content-Disposition": f'attachment; filename="{filename}"'
    }

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers=headers
    )
