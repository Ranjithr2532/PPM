"""
FastAPI Router & Document Generator for ISO Document 068: Engineering Change Note (ECN).
Generates Word (.docx) document matching CMTI-QMS-068/Rev00 specification.
"""

from typing import List, Dict, Any, Optional
import io
from fastapi import APIRouter, HTTPException, status, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from iso.header import add_header_table
from iso.finalfooter import add_footer_table
from iso.sqap import set_cell_shading, set_cell_border, set_cell_margins, add_text

router = APIRouter(prefix="/iso", tags=["ISO Engineering Change Note (Doc 068)"])


# ============================================================
# REQUEST MODELS
# ============================================================

class EngineeringChangeNoteRequest(BaseModel):
    project_title: str = ""
    project_no: str = ""
    customer_name: str = ""
    sub_system: str = ""
    ecn_no: str = ""
    ecn_date: str = ""

    part_description: str = ""

    original_part_no_name: str = ""
    original_rev_no: str = "Rev00"

    changed_part_no_name: str = ""
    changed_rev_no: str = "Rev01"

    reason_for_change: str = ""
    description_of_change: str = ""

    affected_po: bool = False
    affected_po_details: str = ""
    affected_drawing: bool = False
    affected_drawing_details: str = ""
    affected_bom: bool = False
    affected_bom_details: str = ""
    affected_other: bool = False
    affected_other_details: str = ""

    team_member_name: str = ""
    team_member_date: str = ""

    ecn_filled_completely: bool = True
    original_drawing_attached: bool = True
    updated_drawing_attached: bool = True
    due_date_understood: bool = True

    execution_option: str = "A"  # 'A', 'B', 'C', 'D'
    execution_details: str = ""

    prepared_by: str = ""
    reviewed_by: str = ""
    approved_by: str = ""
    group_name: str = ""
    centre_dept: str = ""
    doc_no: str = "068"
    doc_date: str = ""
    filename: str = "ISO_Engineering_Change_Note.docx"


# ============================================================
# DOCUMENT GENERATOR
# ============================================================

def create_engineering_change_note_document(
    project_title: str = "",
    project_no: str = "",
    customer_name: str = "",
    sub_system: str = "",
    ecn_no: str = "",
    ecn_date: str = "",
    part_description: str = "",
    original_part_no_name: str = "",
    original_rev_no: str = "Rev00",
    changed_part_no_name: str = "",
    changed_rev_no: str = "Rev01",
    reason_for_change: str = "",
    description_of_change: str = "",
    affected_po: bool = False,
    affected_po_details: str = "",
    affected_drawing: bool = False,
    affected_drawing_details: str = "",
    affected_bom: bool = False,
    affected_bom_details: str = "",
    affected_other: bool = False,
    affected_other_details: str = "",
    team_member_name: str = "",
    team_member_date: str = "",
    ecn_filled_completely: bool = True,
    original_drawing_attached: bool = True,
    updated_drawing_attached: bool = True,
    due_date_understood: bool = True,
    execution_option: str = "A",
    execution_details: str = "",
    prepared_by: str = "",
    reviewed_by: str = "",
    approved_by: str = "",
    group_name: str = "",
    centre_dept: str = "",
    doc_no: str = "068",
    doc_date: str = ""
) -> Document:
    doc = Document()

    # A4 Portrait margins (0.5" all sides => printable width = 7.27")
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    header_group = (group_name or centre_dept or "SMC").strip().upper()
    if header_group.startswith("G-") or header_group.startswith("C-"):
        header_group = header_group[2:]

    # Add Standard ISO Header Table
    add_header_table(
        section,
        title=f"ENGINEERING CHANGE NOTE-{header_group}" if header_group else "ENGINEERING CHANGE NOTE",
        page_str="1 of 1",
        centre_dept=centre_dept,
        doc_no=doc_no or "068",
        date_str=doc_date or ecn_date
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    border_fmt = {"val": "single", "sz": "4", "color": "000000"}
    hdr_bg = "D9E2EC"

    # --- Project Metadata Card ---
    if project_title or project_no or customer_name or ecn_no:
        meta_table = doc.add_table(rows=1, cols=4)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_table.autofit = False

        col_w = Inches(7.27 / 4.0)
        row = meta_table.rows[0]
        meta_data = [
            ("Project Title:", project_title or "--"),
            ("Project No:", project_no or "--"),
            ("ECN No:", ecn_no or "--"),
            ("Customer:", customer_name or "--")
        ]

        for c_idx, (lbl, val) in enumerate(meta_data):
            cell = row.cells[c_idx]
            cell.width = col_w
            set_cell_border(cell, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
            set_cell_margins(cell, top=20, start=20, bottom=20, end=20)
            set_cell_shading(cell, "F8FAFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_text(p, f"{lbl} ", font_size=8, bold=True)
            add_text(p, str(val), font_size=8, bold=False)

        doc.add_paragraph().paragraph_format.space_after = Pt(3)

    # --- Main Content Table ---
    # We construct a unified structured table matching 068 layout
    content_table = doc.add_table(rows=9, cols=2)
    content_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    content_table.autofit = False

    w_col0 = Inches(5.27)
    w_col1 = Inches(2.0)

    # Helper for cell styling
    def style_cell(cell, width, bg=None):
        cell.width = width
        set_cell_border(cell, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
        set_cell_margins(cell, top=20, start=20, bottom=20, end=20)
        if bg:
            set_cell_shading(cell, bg)

    # Row 0: Part Description (Merged across full width)
    r0 = content_table.rows[0]
    r0_cell = r0.cells[0]
    r0_cell.merge(r0.cells[1])
    style_cell(r0_cell, Inches(7.27), "F4F6F8")
    p0 = r0_cell.paragraphs[0]
    p0.paragraph_format.space_after = Pt(2)
    add_text(p0, "PART DESCRIPTION: ", font_size=9, bold=True)
    add_text(p0, part_description or "----------------", font_size=9, bold=False)

    # Row 1: Original Part / Drawing Name & No + Revision
    r1 = content_table.rows[1]
    style_cell(r1.cells[0], w_col0)
    style_cell(r1.cells[1], w_col1)
    p1_0 = r1.cells[0].paragraphs[0]
    p1_0.paragraph_format.space_after = Pt(0)
    add_text(p1_0, "Original part/drawing name/no: ", font_size=8.5, bold=True)
    add_text(p1_0, original_part_no_name or "--", font_size=8.5)

    p1_1 = r1.cells[1].paragraphs[0]
    p1_1.paragraph_format.space_after = Pt(0)
    add_text(p1_1, "Revision no: ", font_size=8.5, bold=True)
    add_text(p1_1, original_rev_no or "Rev00", font_size=8.5)

    # Row 2: Changed Part / Drawing Name & No + Revision
    r2 = content_table.rows[2]
    style_cell(r2.cells[0], w_col0)
    style_cell(r2.cells[1], w_col1)
    p2_0 = r2.cells[0].paragraphs[0]
    p2_0.paragraph_format.space_after = Pt(0)
    add_text(p2_0, "Changed part/drawing name/no: ", font_size=8.5, bold=True)
    add_text(p2_0, changed_part_no_name or "--", font_size=8.5)

    p2_1 = r2.cells[1].paragraphs[0]
    p2_1.paragraph_format.space_after = Pt(0)
    add_text(p2_1, "Revision no: ", font_size=8.5, bold=True)
    add_text(p2_1, changed_rev_no or "Rev01", font_size=8.5)

    # Row 3: Reason for Change (Merged)
    r3 = content_table.rows[3]
    r3_cell = r3.cells[0]
    r3_cell.merge(r3.cells[1])
    style_cell(r3_cell, Inches(7.27))
    p3 = r3_cell.paragraphs[0]
    p3.paragraph_format.space_after = Pt(2)
    add_text(p3, "REASON FOR CHANGE:\n", font_size=8.5, bold=True)
    add_text(p3, reason_for_change or "N/A", font_size=8.5)

    # Row 4: Description of Change (Merged)
    r4 = content_table.rows[4]
    r4_cell = r4.cells[0]
    r4_cell.merge(r4.cells[1])
    style_cell(r4_cell, Inches(7.27))
    p4 = r4_cell.paragraphs[0]
    p4.paragraph_format.space_after = Pt(2)
    add_text(p4, "DESCRIPTION OF CHANGE:\n", font_size=8.5, bold=True)
    add_text(p4, description_of_change or "N/A", font_size=8.5)

    # Row 5: Affected Documents & Parts (Merged)
    r5 = content_table.rows[5]
    r5_cell = r5.cells[0]
    r5_cell.merge(r5.cells[1])
    style_cell(r5_cell, Inches(7.27))
    p5 = r5_cell.paragraphs[0]
    p5.paragraph_format.space_after = Pt(2)
    add_text(p5, "AFFECTED DOCUMENTS & PARTS:\n", font_size=8.5, bold=True)

    po_chk = "☑" if affected_po else "☐"
    dwg_chk = "☑" if affected_drawing else "☐"
    bom_chk = "☑" if affected_bom else "☐"
    oth_chk = "☑" if affected_other else "☐"

    aff_lines = [
        f"  {po_chk} Purchase order(s)" + (f" - {affected_po_details}" if affected_po_details else ""),
        f"  {dwg_chk} Assembly drawing(s)" + (f" - {affected_drawing_details}" if affected_drawing_details else ""),
        f"  {bom_chk} Bill of materials (BOM)" + (f" - {affected_bom_details}" if affected_bom_details else ""),
        f"  {oth_chk} Other parts (list names)" + (f" - {affected_other_details}" if affected_other_details else "")
    ]
    add_text(p5, "\n".join(aff_lines), font_size=8.5)

    # Row 6: Team Member Details (Prepared By)
    r6 = content_table.rows[6]
    style_cell(r6.cells[0], w_col0)
    style_cell(r6.cells[1], w_col1)
    p6_0 = r6.cells[0].paragraphs[0]
    p6_0.paragraph_format.space_after = Pt(0)
    add_text(p6_0, "Team Member Name: ", font_size=8.5, bold=True)
    add_text(p6_0, team_member_name or prepared_by or "--", font_size=8.5)

    p6_1 = r6.cells[1].paragraphs[0]
    p6_1.paragraph_format.space_after = Pt(0)
    add_text(p6_1, "Date: ", font_size=8.5, bold=True)
    add_text(p6_1, team_member_date or doc_date or "--", font_size=8.5)

    # Row 7: Approval Header (Merged, Header Style)
    r7 = content_table.rows[7]
    r7_cell = r7.cells[0]
    r7_cell.merge(r7.cells[1])
    style_cell(r7_cell, Inches(7.27), hdr_bg)
    p7 = r7_cell.paragraphs[0]
    p7.paragraph_format.space_after = Pt(0)
    add_text(p7, "APPROVAL & REQUIREMENTS REVIEW", font_size=9, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    # Row 8: Requirements Checklist & Execution Strategy (Merged)
    r8 = content_table.rows[8]
    r8_cell = r8.cells[0]
    r8_cell.merge(r8.cells[1])
    style_cell(r8_cell, Inches(7.27))
    p8 = r8_cell.paragraphs[0]
    p8.paragraph_format.space_after = Pt(2)

    # Requirements Table Lines
    q1 = "☑ Yes  ☐ No" if ecn_filled_completely else "☐ Yes  ☑ No"
    q2 = "☑ Yes  ☐ No" if original_drawing_attached else "☐ Yes  ☑ No"
    q3 = "☑ Yes  ☐ No" if updated_drawing_attached else "☐ Yes  ☑ No"
    q4 = "☑ Yes  ☐ No" if due_date_understood else "☐ Yes  ☑ No"

    req_text = (
        "REQUIREMENTS:\n"
        f"  • Is the ECN filled out completely?  ->  {q1}\n"
        f"  • Is the original drawing attached?  ->  {q2}\n"
        f"  • Is the updated drawing attached?   ->  {q3}\n"
        f"  • Is the due date understood by the team?  ->  {q4}\n\n"
        "EXECUTION STRATEGY:\n"
        f"  {'☑' if execution_option == 'A' else '☐'} A - Modify existing part\n"
        f"  {'☑' if execution_option == 'B' else '☐'} B - Create new part\n"
        f"  {'☑' if execution_option == 'C' else '☐'} C - Scrap existing part (if manufacturing already commenced)\n"
        f"  {'☑' if execution_option == 'D' else '☐'} D - Delete existing part (if manufacturing hasn't commenced)"
    )
    if execution_details:
        req_text += f"\n  Execution Notes: {execution_details}"

    add_text(p8, req_text, font_size=8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Standard ISO Footer Table
    doc_code_footer = f"CMTI-{header_group}-QMS-{doc_no or '068'}/Rev00"
    add_footer_table(
        doc,
        prepared_name=prepared_by or team_member_name,
        approved_name=approved_by,
        group_name=header_group,
        doc_code=doc_code_footer,
        in_body=True
    )

    return doc


def generate_engineering_change_note_bytes(
    project_title: str = "",
    project_no: str = "",
    customer_name: str = "",
    sub_system: str = "",
    ecn_no: str = "",
    ecn_date: str = "",
    part_description: str = "",
    original_part_no_name: str = "",
    original_rev_no: str = "Rev00",
    changed_part_no_name: str = "",
    changed_rev_no: str = "Rev01",
    reason_for_change: str = "",
    description_of_change: str = "",
    affected_po: bool = False,
    affected_po_details: str = "",
    affected_drawing: bool = False,
    affected_drawing_details: str = "",
    affected_bom: bool = False,
    affected_bom_details: str = "",
    affected_other: bool = False,
    affected_other_details: str = "",
    team_member_name: str = "",
    team_member_date: str = "",
    ecn_filled_completely: bool = True,
    original_drawing_attached: bool = True,
    updated_drawing_attached: bool = True,
    due_date_understood: bool = True,
    execution_option: str = "A",
    execution_details: str = "",
    prepared_by: str = "",
    reviewed_by: str = "",
    approved_by: str = "",
    group_name: str = "",
    centre_dept: str = "",
    doc_no: str = "068",
    doc_date: str = ""
) -> bytes:
    doc = create_engineering_change_note_document(
        project_title=project_title,
        project_no=project_no,
        customer_name=customer_name,
        sub_system=sub_system,
        ecn_no=ecn_no,
        ecn_date=ecn_date,
        part_description=part_description,
        original_part_no_name=original_part_no_name,
        original_rev_no=original_rev_no,
        changed_part_no_name=changed_part_no_name,
        changed_rev_no=changed_rev_no,
        reason_for_change=reason_for_change,
        description_of_change=description_of_change,
        affected_po=affected_po,
        affected_po_details=affected_po_details,
        affected_drawing=affected_drawing,
        affected_drawing_details=affected_drawing_details,
        affected_bom=affected_bom,
        affected_bom_details=affected_bom_details,
        affected_other=affected_other,
        affected_other_details=affected_other_details,
        team_member_name=team_member_name,
        team_member_date=team_member_date,
        ecn_filled_completely=ecn_filled_completely,
        original_drawing_attached=original_drawing_attached,
        updated_drawing_attached=updated_drawing_attached,
        due_date_understood=due_date_understood,
        execution_option=execution_option,
        execution_details=execution_details,
        prepared_by=prepared_by,
        reviewed_by=reviewed_by,
        approved_by=approved_by,
        group_name=group_name,
        centre_dept=centre_dept,
        doc_no=doc_no,
        doc_date=doc_date
    )
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


# ============================================================
# ROUTER ENDPOINTS
# ============================================================

@router.post("/engineering-change-note/generate")
def generate_engineering_change_note_post(payload: EngineeringChangeNoteRequest):
    try:
        data = generate_engineering_change_note_bytes(
            project_title=payload.project_title,
            project_no=payload.project_no,
            customer_name=payload.customer_name,
            sub_system=payload.sub_system,
            ecn_no=payload.ecn_no,
            ecn_date=payload.ecn_date,
            part_description=payload.part_description,
            original_part_no_name=payload.original_part_no_name,
            original_rev_no=payload.original_rev_no,
            changed_part_no_name=payload.changed_part_no_name,
            changed_rev_no=payload.changed_rev_no,
            reason_for_change=payload.reason_for_change,
            description_of_change=payload.description_of_change,
            affected_po=payload.affected_po,
            affected_po_details=payload.affected_po_details,
            affected_drawing=payload.affected_drawing,
            affected_drawing_details=payload.affected_drawing_details,
            affected_bom=payload.affected_bom,
            affected_bom_details=payload.affected_bom_details,
            affected_other=payload.affected_other,
            affected_other_details=payload.affected_other_details,
            team_member_name=payload.team_member_name,
            team_member_date=payload.team_member_date,
            ecn_filled_completely=payload.ecn_filled_completely,
            original_drawing_attached=payload.original_drawing_attached,
            updated_drawing_attached=payload.updated_drawing_attached,
            due_date_understood=payload.due_date_understood,
            execution_option=payload.execution_option,
            execution_details=payload.execution_details,
            prepared_by=payload.prepared_by,
            reviewed_by=payload.reviewed_by,
            approved_by=payload.approved_by,
            group_name=payload.group_name,
            centre_dept=payload.centre_dept,
            doc_no=payload.doc_no,
            doc_date=payload.doc_date
        )
        fn = payload.filename or f"ISO_Engineering_Change_Note_{payload.ecn_no or payload.project_no or '068'}.docx"
        if not fn.endswith(".docx"):
            fn += ".docx"
        return StreamingResponse(
            io.BytesIO(data),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f'attachment; filename="{fn}"'}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"ECN generation failed: {str(e)}")


@router.get("/engineering-change-note/generate")
def generate_engineering_change_note_get(
    project_title: str = Query("", description="Project Title"),
    project_no: str = Query("", description="Project Number"),
    customer_name: str = Query("", description="Customer Name"),
    ecn_no: str = Query("", description="ECN Number"),
    prepared_by: str = Query("", description="Prepared By"),
    approved_by: str = Query("", description="Approved By"),
    group_name: str = Query("SMC", description="Group Name"),
    centre_dept: str = Query("", description="Centre / Dept"),
    doc_no: str = Query("068", description="Document No"),
    doc_date: str = Query("", description="Document Date"),
    filename: str = Query("ISO_Engineering_Change_Note.docx")
):
    try:
        data = generate_engineering_change_note_bytes(
            project_title=project_title,
            project_no=project_no,
            customer_name=customer_name,
            ecn_no=ecn_no,
            prepared_by=prepared_by,
            approved_by=approved_by,
            group_name=group_name,
            centre_dept=centre_dept,
            doc_no=doc_no,
            doc_date=doc_date
        )
        return StreamingResponse(
            io.BytesIO(data),
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"ECN generation failed: {str(e)}")
