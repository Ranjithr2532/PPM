"""
FastAPI Router & Document Generator for ISO Document 088: Customer Feedback Form.
Generates Word (.docx) document matching CMTI-SMC-QMS-088/Rev00 format.
"""

from typing import List, Dict, Any, Optional
import io
import os
from fastapi import APIRouter, HTTPException, status
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from iso.sqap import set_cell_shading, set_cell_border, set_cell_margins, add_text

router = APIRouter(prefix="/iso", tags=["ISO Customer Feedback (Doc 088)"])


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def set_cell_width(cell, width_inches: float):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(int(width_inches * 1440)))
    tcW.set(qn("w:type"), "dxa")


# ============================================================
# REQUEST MODELS
# ============================================================

class FeedbackRatingItem(BaseModel):
    sl_no: str = ""
    parameter: str = ""
    rating: str = ""


class OfficeUseBlock(BaseModel):
    total: str = "90"
    score_obtained: str = ""
    percentage: str = ""
    conclusion_remarks: str = ""
    action_plan: str = ""


class CustomerFeedbackRequest(BaseModel):
    project_title: str = ""
    project_no: str = ""
    company_name_address: str = ""
    customer_name: str = ""
    doc_no: str = "088"
    doc_date: str = ""
    doc_code: str = "CMTI-SMC-QMS-088/Rev00"
    ratings: Optional[List[FeedbackRatingItem]] = None
    comments: str = ""
    customer_rep_name: str = ""
    office_use: Optional[OfficeUseBlock] = None
    prepared_by: str = ""
    approved_by: str = ""
    group_name: str = "SMC"
    centre_dept: str = "C-SMPM/SMC"
    filename: str = "ISO_Customer_Feedback_088.docx"


DEFAULT_FEEDBACK_PARAMETERS = [
    (1, "Overall, how satisfied were you with the training program?"),
    (2, "Whether the rate quoted for the project is justified?"),
    (3, "Whether required documents / Information submitted on time to adhere to your organizations quality system?"),
    (4, "How do you rate the quality of the product/deliverables? (Product/test reports etc)"),
    (5, "How do you rate the technical Competence of CMTI Team? (Design, Manufacturing, assembly and testing)"),
    (6, "How do you rate the effectiveness of any communication? (Project updates, MOM action points etc.)"),
    (7, "How effectively were the reviews conducted (Design review, Project stage review etc.)"),
    (8, "How was the timeline for delivery maintained? (if project is not completed please rate for stage wise execution)"),
    (9, "Whether all the Technical requirements met?")
]


# ============================================================
# DOCUMENT GENERATOR
# ============================================================

def create_customer_feedback_document(
    project_title: str = "",
    project_no: str = "",
    company_name_address: str = "",
    customer_name: str = "",
    doc_no: str = "088",
    doc_date: str = "",
    doc_code: str = "CMTI-SMC-QMS-088/Rev00",
    ratings: Optional[List[Dict[str, Any]]] = None,
    comments: str = "",
    customer_rep_name: str = "",
    office_use: Optional[Dict[str, Any]] = None,
    prepared_by: str = "",
    approved_by: str = "",
    group_name: str = "SMC",
    centre_dept: str = "C-SMPM/SMC"
) -> Document:
    doc = Document()

    # Setup margins (A4 Portrait, 0.5" all sides => printable width = 7.27")
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    border_fmt = {"val": "single", "sz": "4", "color": "000000"}
    hdr_bg = "F1F5F9"
    full_w = Inches(7.27)

    # 1. Date Header at Top Right
    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_date.paragraph_format.space_before = Pt(0)
    p_date.paragraph_format.space_after = Pt(4)
    run_date_lbl = p_date.add_run("Date: ")
    run_date_lbl.bold = True
    run_date_lbl.font.name = "Arial"
    run_date_lbl.font.size = Pt(10)
    run_date_val = p_date.add_run(doc_date or "")
    run_date_val.font.name = "Arial"
    run_date_val.font.size = Pt(10)

    # 2. Company Info & Instruction Box
    info_table = doc.add_table(rows=1, cols=1)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False

    cell_info = info_table.rows[0].cells[0]
    cell_info.width = full_w
    set_cell_width(cell_info, full_w.inches)
    set_cell_margins(cell_info, top=80, start=100, bottom=80, end=100)
    set_cell_border(cell_info, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)

    p_info = cell_info.paragraphs[0]
    p_info.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_info.paragraph_format.space_after = Pt(2)
    p_info.paragraph_format.line_spacing = 1.15

    # Company Name & Address
    r1 = p_info.add_run("Company Name & Address: ")
    r1.bold = True
    r1.font.name = "Arial"
    r1.font.size = Pt(9.5)
    r1_val = p_info.add_run(f"{company_name_address or customer_name or ''}\n")
    r1_val.font.name = "Arial"
    r1_val.font.size = Pt(9.5)

    # Project Name
    r2 = p_info.add_run("Project Name: ")
    r2.bold = True
    r2.font.name = "Arial"
    r2.font.size = Pt(9.5)
    r2_val = p_info.add_run(f"{project_title or ''}\n\n")
    r2_val.font.name = "Arial"
    r2_val.font.size = Pt(9.5)

    # Scale Instruction
    r3 = p_info.add_run("Please indicate your opinion towards the following parameters on a scale of 1 to 10 (10 being the highest).")
    r3.italic = True
    r3.font.name = "Arial"
    r3.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 3. Parameters & Ratings Table (10 items)
    # Col widths: Sl.No = 0.65", Parameter = 5.62", Rating = 1.0" (sum = 7.27")
    col_w_sl = Inches(0.65)
    col_w_param = Inches(5.62)
    col_w_rate = Inches(1.00)

    # Map supplied ratings by sl_no
    rating_map = {}
    if ratings:
        for r in ratings:
            if isinstance(r, dict):
                rating_map[str(r.get("sl_no"))] = str(r.get("rating") or "")
            elif hasattr(r, "sl_no"):
                rating_map[str(r.sl_no)] = str(r.rating or "")

    table = doc.add_table(rows=11, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_row = table.rows[0]
    hdr_titles = [("Sl.\nNo.", col_w_sl), ("Parameter", col_w_param), ("Rating", col_w_rate)]
    for c_idx, (t, w) in enumerate(hdr_titles):
        cell = hdr_row.cells[c_idx]
        cell.width = w
        set_cell_width(cell, w.inches)
        set_cell_shading(cell, hdr_bg)
        set_cell_margins(cell, top=60, start=60, bottom=60, end=60)
        set_cell_border(cell, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(t)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9.5)

    # Rows 1 to 9: Parameters
    total_score = 0
    rated_count = 0

    for idx, (sl_num, param_text) in enumerate(DEFAULT_FEEDBACK_PARAMETERS):
        row = table.rows[idx + 1]
        rate_val = rating_map.get(str(sl_num), "")

        if rate_val and rate_val.isdigit():
            total_score += int(rate_val)
            rated_count += 1

        # Cell 0: Sl No
        c0 = row.cells[0]
        c0.width = col_w_sl
        set_cell_width(c0, col_w_sl.inches)
        set_cell_margins(c0, top=50, start=40, bottom=50, end=40)
        set_cell_border(c0, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
        add_text(c0, f"{sl_num}.", font_size=9.5, alignment=WD_ALIGN_PARAGRAPH.CENTER)

        # Cell 1: Parameter Description
        c1 = row.cells[1]
        c1.width = col_w_param
        set_cell_width(c1, col_w_param.inches)
        set_cell_margins(c1, top=50, start=60, bottom=50, end=60)
        set_cell_border(c1, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
        add_text(c1, param_text, font_size=9.5, alignment=WD_ALIGN_PARAGRAPH.LEFT)

        # Cell 2: Rating
        c2 = row.cells[2]
        c2.width = col_w_rate
        set_cell_width(c2, col_w_rate.inches)
        set_cell_margins(c2, top=50, start=40, bottom=50, end=40)
        set_cell_border(c2, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
        add_text(c2, rate_val, font_size=9.5, bold=True, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    # Row 10: Comments / Complaints / Suggestions
    row10 = table.rows[10]
    c10_sl = row10.cells[0]
    c10_sl.width = col_w_sl
    set_cell_width(c10_sl, col_w_sl.inches)
    set_cell_margins(c10_sl, top=50, start=40, bottom=50, end=40)
    set_cell_border(c10_sl, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
    add_text(c10_sl, "10.", font_size=9.5, alignment=WD_ALIGN_PARAGRAPH.CENTER)

    c10_content = row10.cells[1]
    c10_content.merge(row10.cells[2])
    set_cell_width(c10_content, col_w_param.inches + col_w_rate.inches)
    set_cell_margins(c10_content, top=50, start=60, bottom=50, end=60)
    set_cell_border(c10_content, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)

    p10 = c10_content.paragraphs[0]
    p10.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r10_lbl = p10.add_run("Your  Complaints / Suggestion / Comments:\n")
    r10_lbl.bold = True
    r10_lbl.font.name = "Arial"
    r10_lbl.font.size = Pt(9.5)
    r10_val = p10.add_run(comments or "")
    r10_val.font.name = "Arial"
    r10_val.font.size = Pt(9.5)

    # 4. Customer Seal & Representative Block
    seal_table = doc.add_table(rows=1, cols=2)
    seal_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    seal_table.autofit = False

    half_w = Inches(7.27 / 2.0)
    cell_seal = seal_table.rows[0].cells[0]
    cell_rep = seal_table.rows[0].cells[1]

    cell_seal.width = half_w
    cell_rep.width = half_w
    set_cell_width(cell_seal, half_w.inches)
    set_cell_width(cell_rep, half_w.inches)
    set_cell_margins(cell_seal, top=50, start=60, bottom=50, end=60)
    set_cell_margins(cell_rep, top=50, start=60, bottom=50, end=60)
    set_cell_border(cell_seal, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
    set_cell_border(cell_rep, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)

    p_seal = cell_seal.paragraphs[0]
    p_seal.paragraph_format.line_spacing = 1.15
    r_seal = p_seal.add_run("Company Seal\n\n")
    r_seal.bold = True
    r_seal.font.name = "Arial"
    r_seal.font.size = Pt(9.5)

    p_rep = cell_rep.paragraphs[0]
    p_rep.paragraph_format.line_spacing = 1.15
    r_rep = p_rep.add_run("Customer representative Name & Sign:\n")
    r_rep.bold = True
    r_rep.font.name = "Arial"
    r_rep.font.size = Pt(9.5)
    if customer_rep_name:
        r_rep_val = p_rep.add_run(f"{customer_rep_name}\n")
        r_rep_val.font.name = "Arial"
        r_rep_val.font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # 5. FOR CMTI OFFICE USE ONLY
    p_office_hdr = doc.add_paragraph()
    p_office_hdr.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_office_hdr.paragraph_format.space_before = Pt(2)
    p_office_hdr.paragraph_format.space_after = Pt(2)
    r_off = p_office_hdr.add_run("FOR CMTI OFFICE USE ONLY.")
    r_off.bold = True
    r_off.font.name = "Arial"
    r_off.font.size = Pt(9.5)

    # Office use summary table (5 cols)
    # Col widths: Total = 1.0", Score Obtained = 1.4", % Of Rating = 1.2", Conclusion/Remarks = 2.0", Action plan = 1.67"
    off_widths = [Inches(1.00), Inches(1.40), Inches(1.20), Inches(2.00), Inches(1.67)]
    off_headers = ["Total", "Score Obtained", "% Of Rating", "Conclusion/\nRemarks", "Action plan"]

    ou = office_use or {}
    total_val = ou.get("total") or "90"
    score_val = ou.get("score_obtained") or (str(total_score) if rated_count > 0 else "")
    pct_val = ou.get("percentage") or (f"{(total_score / 90.0) * 100:.2f}%" if rated_count > 0 else "")
    concl_val = ou.get("conclusion_remarks") or ""
    action_val = ou.get("action_plan") or ""

    off_table = doc.add_table(rows=2, cols=5)
    off_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    off_table.autofit = False

    # Header
    off_hdr_row = off_table.rows[0]
    for c_idx, (t, w) in enumerate(zip(off_headers, off_widths)):
        cell = off_hdr_row.cells[c_idx]
        cell.width = w
        set_cell_width(cell, w.inches)
        set_cell_shading(cell, hdr_bg)
        set_cell_margins(cell, top=50, start=40, bottom=50, end=40)
        set_cell_border(cell, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(t)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9.0)

    # Values
    off_val_row = off_table.rows[1]
    val_list = [total_val, score_val, pct_val, concl_val, action_val]
    for c_idx, (v, w) in enumerate(zip(val_list, off_widths)):
        cell = off_val_row.cells[c_idx]
        cell.width = w
        set_cell_width(cell, w.inches)
        set_cell_margins(cell, top=50, start=40, bottom=50, end=40)
        set_cell_border(cell, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        align = WD_ALIGN_PARAGRAPH.CENTER if c_idx < 3 else WD_ALIGN_PARAGRAPH.LEFT
        add_text(cell, v, font_size=9.0, bold=(c_idx < 3), alignment=align)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 6. Institute Address Box Footer
    inst_table = doc.add_table(rows=1, cols=1)
    inst_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    inst_table.autofit = False

    cell_inst = inst_table.rows[0].cells[0]
    cell_inst.width = full_w
    set_cell_width(cell_inst, full_w.inches)
    set_cell_margins(cell_inst, top=50, start=60, bottom=50, end=60)
    set_cell_border(cell_inst, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)

    p_inst = cell_inst.paragraphs[0]
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.line_spacing = 1.15
    r_i1 = p_inst.add_run("CENTRAL MANUFACTURING TECHNOLOGY INSTITUTE\n")
    r_i1.bold = True
    r_i1.font.name = "Arial"
    r_i1.font.size = Pt(9.5)
    r_i2 = p_inst.add_run("TUMKUR ROAD, BANGALORE 560 022")
    r_i2.bold = True
    r_i2.font.name = "Arial"
    r_i2.font.size = Pt(9.5)

    # Format code bottom left
    p_code = doc.add_paragraph()
    p_code.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_code.paragraph_format.space_before = Pt(2)
    r_code = p_code.add_run(f"Format No : {doc_code or 'CMTI-SMC-QMS-088/Rev00'}")
    r_code.font.name = "Arial"
    r_code.font.size = Pt(8.5)
    r_code.font.color.rgb = RGBColor(100, 100, 100)

    return doc


# ============================================================
# FASTAPI GENERATE ENDPOINT
# ============================================================

@router.post("/customer-feedback/generate")
def generate_customer_feedback(req: CustomerFeedbackRequest):
    try:
        raw_ratings = [r.dict() if hasattr(r, 'dict') else r for r in (req.ratings or [])]
        raw_office_use = req.office_use.dict() if req.office_use and hasattr(req.office_use, 'dict') else (req.office_use or {})

        doc = create_customer_feedback_document(
            project_title=req.project_title,
            project_no=req.project_no,
            company_name_address=req.company_name_address,
            customer_name=req.customer_name,
            doc_no=req.doc_no or "088",
            doc_date=req.doc_date,
            doc_code=req.doc_code or "CMTI-SMC-QMS-088/Rev00",
            ratings=raw_ratings,
            comments=req.comments,
            customer_rep_name=req.customer_rep_name,
            office_use=raw_office_use,
            prepared_by=req.prepared_by,
            approved_by=req.approved_by,
            group_name=req.group_name or "SMC",
            centre_dept=req.centre_dept or "C-SMPM/SMC"
        )

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        filename = req.filename or f"ISO_Customer_Feedback_{req.doc_no or '088'}.docx"
        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'}
        )
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to generate Customer Feedback (.docx): {str(e)}"
        )
