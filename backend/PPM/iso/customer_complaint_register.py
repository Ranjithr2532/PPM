"""
FastAPI Router & Document Generator for ISO Document 086: Customer Complaint Register.
Generates Word (.docx) document matching CMTI-QMS-SMC-086/Rev00 format.
"""

from typing import List, Dict, Any, Optional
import io
import os
from fastapi import APIRouter, HTTPException, status
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from iso.header import add_header_table, normalize_centre_dept
from iso.finalfooter import add_footer_table
from iso.sqap import set_cell_shading, set_cell_border, set_cell_margins, add_text

def set_cell_width(cell, width_inches: float):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(int(width_inches * 1440)))
    tcW.set(qn("w:type"), "dxa")

router = APIRouter(prefix="/iso", tags=["ISO Customer Complaint Register (Doc 086)"])


# ============================================================
# REQUEST MODELS
# ============================================================

class ComplaintRowRequest(BaseModel):
    sl_no: str = ""
    date: str = ""
    complaint_ref: str = ""
    complaint_desc: str = ""
    customer_details: str = ""
    nc_no: str = ""
    remarks: str = ""


class CustomerComplaintRegisterRequest(BaseModel):
    project_title: str = ""
    project_no: str = ""
    customer_name: str = ""
    doc_no: str = "086"
    doc_date: str = ""
    doc_code: str = "CMTI-QMS-SMC-086/Rev00"
    centre_dept: str = "C-SMPM/SMC"
    group_name: str = "SMC"
    page_str: str = "1 of 1"
    rows: Optional[List[ComplaintRowRequest]] = None
    prepared_by: str = ""
    approved_by: str = ""
    filename: str = "ISO_Customer_Complaint_Register_086.docx"


DEFAULT_COMPLAINT_HEADERS = [
    "Sl.No.",
    "Date",
    "Complaint\nreference",
    "Complaint description",
    "Customer\nName & Address of\nCustomer",
    "NC No. as per NC\nregister",
    "Remarks"
]


# ============================================================
# DOCUMENT GENERATOR
# ============================================================

def create_customer_complaint_register_document(
    project_title: str = "",
    project_no: str = "",
    customer_name: str = "",
    doc_no: str = "086",
    doc_date: str = "",
    doc_code: str = "CMTI-QMS-SMC-086/Rev00",
    centre_dept: str = "C-SMPM/SMC",
    group_name: str = "SMC",
    page_str: str = "1 of 1",
    rows: Optional[List[Dict[str, Any]]] = None,
    prepared_by: str = "",
    approved_by: str = ""
) -> Document:
    doc = Document()

    # Landscape A4 (11.69" x 8.27", 0.5" margins => 10.69" printable width)
    for section in doc.sections:
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    header_group = (group_name or centre_dept or "SMC").strip().upper()
    if header_group.startswith("G-") or header_group.startswith("C-"):
        header_group = header_group[2:]

    # Add Standard Header Table
    add_header_table(
        doc.sections[0],
        title=f"CUSTOMER COMPLAINT REGISTER",
        page_str=page_str or "1 of 1",
        centre_dept=centre_dept or "C-SMPM/SMC",
        doc_no=doc_no or "086",
        date_str=doc_date or ""
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    border_fmt = {"val": "single", "sz": "4", "color": "000000"}
    hdr_bg = "F1F5F9"

    # 7 Columns Table widths (summing to ~10.69 inches)
    col_widths = [
        Inches(0.80),   # Sl.No.
        Inches(1.05),   # Date
        Inches(1.60),   # Complaint reference
        Inches(2.40),   # Complaint description
        Inches(2.40),   # Customer Name & Address
        Inches(1.24),   # NC No.
        Inches(1.20)    # Remarks
    ]

    raw_rows = rows or []
    row_count = max(len(raw_rows), 7)  # At least 7 rows as in the reference template

    table = doc.add_table(rows=row_count + 1, cols=7)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_row = table.rows[0]
    for c_idx, title in enumerate(DEFAULT_COMPLAINT_HEADERS):
        cell = hdr_row.cells[c_idx]
        cell.width = col_widths[c_idx]
        set_cell_width(cell, col_widths[c_idx].inches)
        set_cell_shading(cell, hdr_bg)
        set_cell_margins(cell, top=80, start=60, bottom=80, end=60)
        set_cell_border(cell, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(title)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(10)

    # Data Rows
    for r_idx in range(row_count):
        row_data = raw_rows[r_idx] if r_idx < len(raw_rows) else {}
        if hasattr(row_data, 'dict'):
            row_data = row_data.dict()

        sl = str(row_data.get("sl_no") or (r_idx + 1 if r_idx < len(raw_rows) else ""))
        dt = str(row_data.get("date") or "")
        ref = str(row_data.get("complaint_ref") or "")
        desc = str(row_data.get("complaint_desc") or "")
        cust = str(row_data.get("customer_details") or "")
        nc = str(row_data.get("nc_no") or "")
        rem = str(row_data.get("remarks") or "")

        cell_values = [sl, dt, ref, desc, cust, nc, rem]
        row_cells = table.rows[r_idx + 1].cells

        for c_idx, val in enumerate(cell_values):
            cell = row_cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_width(cell, col_widths[c_idx].inches)
            set_cell_margins(cell, top=80, start=60, bottom=80, end=60)
            set_cell_border(cell, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            align = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 1, 5] else WD_ALIGN_PARAGRAPH.LEFT
            add_text(cell, val, font_size=9.5, alignment=align)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Footer Table (Prepared By, Approved By, Institute Name & Address, Document Code)
    footer_table = doc.add_table(rows=3, cols=2)
    footer_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    footer_table.autofit = False

    half_w = Inches(10.69 / 2.0)

    # Row 0: Prepared by / Approved By
    row0 = footer_table.rows[0]
    cell_p = row0.cells[0]
    cell_a = row0.cells[1]

    cell_p.width = half_w
    cell_a.width = half_w
    set_cell_width(cell_p, half_w.inches)
    set_cell_width(cell_a, half_w.inches)
    set_cell_margins(cell_p, top=60, start=80, bottom=60, end=80)
    set_cell_margins(cell_a, top=60, start=80, bottom=60, end=80)
    set_cell_border(cell_p, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)
    set_cell_border(cell_a, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)

    # Prepared by content
    p_prep = cell_p.paragraphs[0]
    p_prep.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_prep.paragraph_format.line_spacing = 1.15
    run = p_prep.add_run("Prepared by\n")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(9.5)
    r_name = p_prep.add_run(f"Name: {prepared_by or ''}\n")
    r_name.font.name = "Arial"
    r_name.font.size = Pt(9.5)
    r_sig = p_prep.add_run("Signature:")
    r_sig.font.name = "Arial"
    r_sig.font.size = Pt(9.5)

    # Approved By (MR) content
    p_app = cell_a.paragraphs[0]
    p_app.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_app.paragraph_format.line_spacing = 1.15
    run = p_app.add_run("Approved By (MR)\n")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(9.5)
    r_name = p_app.add_run(f"Name: {approved_by or ''}\n")
    r_name.font.name = "Arial"
    r_name.font.size = Pt(9.5)
    r_sig = p_app.add_run("Signature:")
    r_sig.font.name = "Arial"
    r_sig.font.size = Pt(9.5)

    # Row 1: Institute Name & Address across full width
    row1 = footer_table.rows[1]
    cell_inst = row1.cells[0]
    cell_inst.merge(row1.cells[1])
    set_cell_width(cell_inst, Inches(10.69).inches)
    set_cell_margins(cell_inst, top=60, start=80, bottom=60, end=80)
    set_cell_border(cell_inst, top=border_fmt, bottom=border_fmt, left=border_fmt, right=border_fmt)

    p_inst = cell_inst.paragraphs[0]
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.line_spacing = 1.1
    r_inst1 = p_inst.add_run("CENTRAL MANUFACTURING TECHNOLOGY INSTITUTE\n")
    r_inst1.bold = True
    r_inst1.font.name = "Arial"
    r_inst1.font.size = Pt(9.5)
    r_inst2 = p_inst.add_run("TUMKUR ROAD, BANGALORE 560 022")
    r_inst2.bold = True
    r_inst2.font.name = "Arial"
    r_inst2.font.size = Pt(9.5)

    # Row 2: Bottom Code (CMTI-QMS-SMC-086/Rev00)
    row2 = footer_table.rows[2]
    cell_code = row2.cells[0]
    cell_code.merge(row2.cells[1])
    set_cell_width(cell_code, Inches(10.69).inches)
    set_cell_margins(cell_code, top=40, start=80, bottom=40, end=80)
    # No top border for clean bottom line look
    no_border = {"val": "none"}
    set_cell_border(cell_code, top=no_border, bottom=no_border, left=no_border, right=no_border)

    p_code = cell_code.paragraphs[0]
    p_code.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_c = p_code.add_run(doc_code or f"CMTI-QMS-{header_group}-086/Rev00")
    r_c.font.name = "Arial"
    r_c.font.size = Pt(8.5)
    r_c.font.color.rgb = RGBColor(100, 100, 100)

    return doc


# ============================================================
# FASTAPI GENERATE ENDPOINT
# ============================================================

@router.post("/customer-complaint-register/generate")
def generate_customer_complaint_register(req: CustomerComplaintRegisterRequest):
    try:
        raw_rows = [r.dict() if hasattr(r, 'dict') else r for r in (req.rows or [])]

        doc = create_customer_complaint_register_document(
            project_title=req.project_title,
            project_no=req.project_no,
            customer_name=req.customer_name,
            doc_no=req.doc_no or "086",
            doc_date=req.doc_date,
            doc_code=req.doc_code or "CMTI-QMS-SMC-086/Rev00",
            centre_dept=req.centre_dept or "C-SMPM/SMC",
            group_name=req.group_name or "SMC",
            page_str=req.page_str or "1 of 1",
            rows=raw_rows,
            prepared_by=req.prepared_by,
            approved_by=req.approved_by
        )

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        filename = req.filename or f"ISO_Customer_Complaint_Register_{req.doc_no or '086'}.docx"
        return StreamingResponse(
            buffer,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'}
        )
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Failed to generate Customer Complaint Register (.docx): {str(e)}"
        )
